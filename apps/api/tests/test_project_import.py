from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook

ADMIN_HEADERS = {"X-Test-Session": "active-admin"}
HEADER_ROW = ["STT", "Ma de tai", "Ma nhom", "Ten de tai Tieng Anh/ Tieng Nhat", "Ten de tai Tieng Viet", "GVHD", "GVHD2"]


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER_ROW)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, semester_id: int, content: bytes, headers: dict[str, str] = ADMIN_HEADERS):
    return client.post(
        f"/api/v1/projects/import?semesterId={semester_id}",
        headers=headers,
        files={"file": ("projects.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def _seed_semester_and_lecturer(client) -> tuple[int, str]:
    """seed-fixture gives a clean ACTIVE semester + a couple of lecturers to import against."""
    client.post("/api/v1/admin/seed-fixture", headers=ADMIN_HEADERS)
    semester_id = client.get("/api/v1/semesters", headers=ADMIN_HEADERS).json()["data"][0]["id"]
    suffix = uuid4().hex[:8]
    lecturer_code = f"GVIMP{suffix}"
    create = client.post(
        "/api/v1/accounts",
        headers=ADMIN_HEADERS,
        json={
            "email": f"gvimp.{suffix}@example.com",
            "displayName": "Test GVHD",
            "password": "P@ssword12345",
            "role": "LECTURER",
            "lecturerCode": lecturer_code,
        },
    )
    assert create.status_code == 201, create.text
    return semester_id, lecturer_code


@pytest.mark.integration
def test_import_projects_creates_project_group_and_main_supervisor(client):
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    project_code = f"PRJ{suffix}"
    group_code = f"GRP{suffix}"
    content = _workbook([[1, project_code, group_code, "Title EN", "Title VI", lecturer_code.lower(), None]])

    response = _upload(client, semester_id, content)

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["created"] == 1
    assert body["updated"] == 0
    assert body["skipped"] == 0
    assert body["errors"] == []


@pytest.mark.integration
def test_import_projects_upserts_on_rerun_instead_of_duplicating(client):
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    project_code = f"PRJ{suffix}"
    group_code = f"GRP{suffix}"
    content = _workbook([[1, project_code, group_code, "Title EN", "Title VI", lecturer_code, None]])

    first = _upload(client, semester_id, content)
    assert first.json()["created"] == 1

    second = _upload(client, semester_id, content)
    assert second.status_code == 201, second.text
    body = second.json()
    assert body["created"] == 0
    assert body["updated"] == 1
    assert body["skipped"] == 0


@pytest.mark.integration
def test_import_projects_row_with_unmatched_gvhd_is_skipped_not_fatal(client):
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    good_code = f"PRJ{suffix}A"
    bad_code = f"PRJ{suffix}B"
    content = _workbook(
        [
            [1, good_code, f"GRP{suffix}A", "Title EN", "Title VI", lecturer_code, None],
            [2, bad_code, f"GRP{suffix}B", "Title EN", "Title VI", "NO-SUCH-CODE", None],
        ]
    )

    response = _upload(client, semester_id, content)

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["created"] == 1
    assert body["skipped"] == 1
    assert body["errors"][0]["code"] == "GVHD_NOT_FOUND"


@pytest.mark.integration
def test_import_projects_finds_header_row_below_a_banner_row(client):
    """Real user file (DanhSachDeTai_FA26.xlsx) has a title banner in row 1 and the real
    header in row 2 — the naive "row 0 is always the header" assumption silently misparsed
    every single row. Regression test for that exact shape."""
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    project_code = f"PRJ{suffix}"
    group_code = f"GRP{suffix}"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["PROJECTS INFORMATION"])  # banner row, not the header
    sheet.append(HEADER_ROW)
    sheet.append([1, project_code, group_code, "Title EN", "Title VI", lecturer_code, None])
    buffer = BytesIO()
    workbook.save(buffer)

    response = _upload(client, semester_id, buffer.getvalue())

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["created"] == 1
    assert body["skipped"] == 0


HEADER_ROW_WITH_DEPARTMENT = ["STT", "Ma de tai", "Ma nhom", "Ten de tai Tieng Anh/ Tieng Nhat", "Ten de tai Tieng Viet", "Department", "GVHD", "GVHD2"]
# Fixed, reused across runs — the major upsert (ON CONFLICT DO UPDATE) makes this idempotent,
# so this never leaks a new majors row per test run the way a per-run uuid code would.
_TEST_MAJOR_CODE = "ZZTEST"


def _project_row(code: str, group: str, dept, lecturer_code: str) -> list[object]:
    return [1, code, group, "Title EN", "Title VI", dept, lecturer_code, None]


def _get_major_code(client, semester_id: int, project_code: str) -> str | None:
    projects = client.get("/api/v1/projects", headers=ADMIN_HEADERS, params={"semesterId": semester_id}).json()
    return next((p.get("majorCode") for p in projects if p["code"] == project_code.upper()), None)


@pytest.mark.integration
def test_import_projects_maps_department_to_major_and_creates_new_ones(client):
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER_ROW_WITH_DEPARTMENT)
    sheet.append(_project_row(f"PRJ{suffix}A", f"GRP{suffix}A", "SE", lecturer_code))
    sheet.append(_project_row(f"PRJ{suffix}B", f"GRP{suffix}B", _TEST_MAJOR_CODE, lecturer_code))
    buffer = BytesIO()
    workbook.save(buffer)

    response = _upload(client, semester_id, buffer.getvalue())

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["created"] == 2
    assert body["skipped"] == 0
    assert _get_major_code(client, semester_id, f"PRJ{suffix}A") == "SE"
    assert _get_major_code(client, semester_id, f"PRJ{suffix}B") == _TEST_MAJOR_CODE


@pytest.mark.integration
def test_import_projects_rerun_with_changed_department_updates_major(client):
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    code, group = f"PRJ{suffix}", f"GRP{suffix}"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER_ROW_WITH_DEPARTMENT)
    sheet.append(_project_row(code, group, "SE", lecturer_code))
    buffer = BytesIO()
    workbook.save(buffer)
    first = _upload(client, semester_id, buffer.getvalue())
    assert first.json()["created"] == 1
    assert _get_major_code(client, semester_id, code) == "SE"

    workbook2 = Workbook()
    sheet2 = workbook2.active
    sheet2.append(HEADER_ROW_WITH_DEPARTMENT)
    sheet2.append(_project_row(code, group, _TEST_MAJOR_CODE, lecturer_code))
    buffer2 = BytesIO()
    workbook2.save(buffer2)
    second = _upload(client, semester_id, buffer2.getvalue())

    assert second.status_code == 201, second.text
    assert second.json()["updated"] == 1
    assert _get_major_code(client, semester_id, code) == _TEST_MAJOR_CODE


@pytest.mark.integration
def test_import_projects_rerun_with_blank_department_keeps_existing_major(client):
    """The FE's own documented template has no Department column at all — a sheet built from
    it must NOT reset an existing project's (possibly non-SE) major back to the SE default."""
    semester_id, lecturer_code = _seed_semester_and_lecturer(client)
    suffix = uuid4().hex[:8]
    code, group = f"PRJ{suffix}", f"GRP{suffix}"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER_ROW_WITH_DEPARTMENT)
    sheet.append(_project_row(code, group, _TEST_MAJOR_CODE, lecturer_code))
    buffer = BytesIO()
    workbook.save(buffer)
    first = _upload(client, semester_id, buffer.getvalue())
    assert first.json()["created"] == 1
    assert _get_major_code(client, semester_id, code) == _TEST_MAJOR_CODE

    # Re-import using the plain template (no Department column at all).
    content = _workbook([[1, code, group, "Title EN", "Title VI 2", lecturer_code, None]])
    second = _upload(client, semester_id, content)

    assert second.status_code == 201, second.text
    assert second.json()["updated"] == 1
    assert _get_major_code(client, semester_id, code) == _TEST_MAJOR_CODE


@pytest.mark.integration
def test_import_projects_no_header_row_returns_422(client):
    semester_id, _ = _seed_semester_and_lecturer(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Not", "A", "Header", "Row"])
    sheet.append(["just", "some", "data", "here"])
    buffer = BytesIO()
    workbook.save(buffer)

    response = _upload(client, semester_id, buffer.getvalue())

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "IMPORT_INVALID_FILE"


@pytest.mark.integration
def test_import_projects_requires_admin_or_manager_role(client):
    content = _workbook([[1, "PRJX001", "GRPX001", "Title EN", "Title VI", "SOMEONE", None]])
    response = client.post(
        "/api/v1/projects/import?semesterId=1",
        headers={"X-Test-Session": "active-student"},
        files={"file": ("projects.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 403
